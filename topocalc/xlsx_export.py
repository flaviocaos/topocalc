"""
xlsx_export.py
===============
Gera a planilha .xlsx com as abas:

- Pontos_Resumo: valores brutos (entrada) + graus/DMS/deltas/distância
  plana calculados por FÓRMULA nativa do Excel (não hardcoded). UTM e
  distância geodésica entram como VALORES DE REFERÊNCIA calculados
  externamente (pyproj/Karney) e claramente documentados como tal, porque
  reproduzir a projeção UTM completa e a solução geodésica elipsoidal em
  fórmulas de planilha nativas não é praticável nem confiável.
- Metadados_Arquivos
- Auditoria_Bytes
- Atualização_Coord_Oficial: input de coordenada oficial da base + fórmulas
  que recalculam os pontos finais por translação dos deltas.
- Metodologia_Limitações
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

from .pipeline import ProcessingResult
from .assets import logo_png_path

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
NOTE_FONT = Font(italic=True, color="666666")


def _style_header(ws, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize(ws, ncols: int, min_w: int = 12, max_w: int = 42) -> None:
    for c in range(1, ncols + 1):
        col_letter = get_column_letter(c)
        max_len = min_w
        for cell in ws[col_letter]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)) + 2)
        ws.column_dimensions[col_letter].width = min(max_len, max_w)


def _sheet_capa(wb, result: ProcessingResult) -> None:
    ws = wb.active
    ws.title = "Capa"
    ws.sheet_view.showGridLines = False

    logo_path = logo_png_path()
    if logo_path is not None:
        img = XLImage(str(logo_path))
        img.width = img.width * 0.28
        img.height = img.height * 0.28
        ws.add_image(img, "B2")

    ws.cell(row=14, column=2, value="Pacote Geodésico — TopoCalc")
    ws.cell(row=14, column=2).font = Font(bold=True, size=16, color="3A3A39")

    ws.cell(row=16, column=2, value=f"Gerado em: {result.base_file.header_date or 'n/d'}")
    ws.cell(row=16, column=2).font = Font(color="9A9A98")

    ws.cell(row=18, column=2, value=(
        "Este arquivo consolida a aba Pontos_Resumo (coordenadas, deltas e distâncias), "
        "Metadados_Arquivos, Auditoria_Bytes, Atualização_Coord_Oficial e Metodologia_Limitações."
    ))
    ws.cell(row=18, column=2).alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=18, start_column=2, end_row=20, end_column=8)
    ws.column_dimensions["B"].width = 14
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 12


def _sheet_pontos_resumo(wb, result: ProcessingResult) -> None:
    ws = wb.create_sheet("Pontos_Resumo")

    headers = [
        "Ponto", "Tipo", "Lat_rad (entrada)", "Lon_rad (entrada)", "Alt_m (entrada)",
        "Lat_graus (fórmula)", "Lon_graus (fórmula)",
        "Lat_DMS (fórmula)", "Lon_DMS (fórmula)",
        "UTM_Zona (ref. externa)", "UTM_Este_m (ref. externa)", "UTM_Norte_m (ref. externa)",
        "Delta_E_m (fórmula)", "Delta_N_m (fórmula)", "Delta_h_m (fórmula)",
        "Dist_Plana_m (fórmula)", "Dist_Geodesica_m (ref. externa)", "Azimute_graus (ref. externa)",
    ]
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    base = result.base_geo
    all_points = [base] + result.points_geo

    # Localiza a linha da base para referenciar em fórmulas de delta.
    base_row = 2

    for i, p in enumerate(all_points):
        row = i + 2
        tipo = "BASE" if p is base else "PONTO"
        # A/B: identificação
        ws.cell(row=row, column=1, value=p.label)
        ws.cell(row=row, column=2, value=tipo)
        # C/D/E: entrada crua
        import math
        ws.cell(row=row, column=3, value=math.radians(p.lat_deg))
        ws.cell(row=row, column=4, value=math.radians(p.lon_deg))
        ws.cell(row=row, column=5, value=p.alt_m)

        # F/G: graus decimais POR FÓRMULA a partir de C/D (radianos -> graus)
        ws.cell(row=row, column=6, value=f"=DEGREES(C{row})")
        ws.cell(row=row, column=7, value=f"=DEGREES(D{row})")

        # H/I: DMS por fórmula, robusto a valores negativos (usa ABS() e
        # aplica o sinal/hemisfério separadamente — mesma lógica do fix
        # aplicado em geo_calc.deg_to_dms).
        lat_dms_formula = (
            f'=IF(F{row}>=0,"N","S")&" "&'
            f'TEXT(INT(ABS(F{row})),"00")&"°"&'
            f'TEXT(INT((ABS(F{row})-INT(ABS(F{row})))*60),"00")&"\'"&'
            f'TEXT(((ABS(F{row})-INT(ABS(F{row})))*60-INT((ABS(F{row})-INT(ABS(F{row})))*60))*60,"00.000")&CHAR(34)'
        )
        lon_dms_formula = (
            f'=IF(G{row}>=0,"E","W")&" "&'
            f'TEXT(INT(ABS(G{row})),"00")&"°"&'
            f'TEXT(INT((ABS(G{row})-INT(ABS(G{row})))*60),"00")&"\'"&'
            f'TEXT(((ABS(G{row})-INT(ABS(G{row})))*60-INT((ABS(G{row})-INT(ABS(G{row})))*60))*60,"00.000")&CHAR(34)'
        )
        ws.cell(row=row, column=8, value=lat_dms_formula)
        ws.cell(row=row, column=9, value=lon_dms_formula)

        # J/K/L: UTM como valor de referência externo (documentado)
        ws.cell(row=row, column=10, value=f"{p.utm_zone}{p.utm_hemisphere}")
        ws.cell(row=row, column=11, value=p.utm_easting)
        ws.cell(row=row, column=12, value=p.utm_northing)

        if p is base:
            for col in (13, 14, 15, 16, 17, 18):
                ws.cell(row=row, column=col, value="—")
        else:
            # M/N/O: deltas por fórmula em relação à linha da base
            ws.cell(row=row, column=13, value=f"=K{row}-K{base_row}")
            ws.cell(row=row, column=14, value=f"=L{row}-L{base_row}")
            ws.cell(row=row, column=15, value=f"=E{row}-E{base_row}")
            # P: distância plana por fórmula (Pitágoras) a partir dos deltas
            ws.cell(row=row, column=16, value=f"=SQRT(M{row}^2+N{row}^2)")
            # Q/R: distância geodésica e azimute como referência externa
            ws.cell(row=row, column=17, value=p.dist_geodesica_m)
            ws.cell(row=row, column=18, value=p.azimute_deg)

    note_row = len(all_points) + 3
    nota_texto = (
        "Nota: colunas marcadas '(fórmula)' são recalculadas nativamente pelo Excel a partir "
        "das colunas de entrada. Colunas marcadas '(ref. externa)' (UTM, distância geodésica, "
        "azimute) foram calculadas fora da planilha com pyproj/PROJ (elipsoide WGS84, solução "
        "geodésica de Karney) e são inseridas como valores de referência, não fórmulas — "
        "reproduzir a projeção UTM e a solução geodésica elipsoidal em fórmulas nativas do "
        "Excel não é praticável nem confiável."
    )
    if result.cross_zone_points:
        pontos_str = ", ".join(p.label for p in result.cross_zone_points)
        nota_texto += (
            f" Atenção: o(s) ponto(s) {pontos_str} está(ão) num fuso UTM diferente do fuso da "
            f"BASE e foi(foram) reprojetado(s) para o fuso da BASE (coluna UTM_Zona) para permitir "
            f"o cálculo de Delta_E/Delta_N/Dist_Plana num único sistema de coordenadas — os "
            f"valores de UTM_Este_m/UTM_Norte_m destes pontos NÃO são as coordenadas nativas do "
            f"fuso onde o ponto realmente está, e sim a versão reprojetada para o fuso da BASE."
        )
    ws.cell(row=note_row, column=1, value=nota_texto)
    ws.cell(row=note_row, column=1).font = NOTE_FONT
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=len(headers))
    ws.cell(row=note_row, column=1).alignment = Alignment(wrap_text=True)

    for col in (3, 4):
        for row in range(2, len(all_points) + 2):
            ws.cell(row=row, column=col).number_format = "0.00000000"
    for col in (6, 7):
        for row in range(2, len(all_points) + 2):
            ws.cell(row=row, column=col).number_format = "0.00000000"

    _autosize(ws, len(headers))
    ws.freeze_panes = "A2"


def _sheet_metadados(wb, result: ProcessingResult) -> None:
    ws = wb.create_sheet("Metadados_Arquivos")
    headers = [
        "Arquivo", "Ponto", "Tipo", "Versão formato", "Data cabeçalho",
        "Modelo receptor", "Firmware", "Serial/ID", "Contagem ZDA",
        "Sessão início (UTC)", "Sessão fim (UTC)",
    ]
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    for f in result.files:
        ws.append([
            f.path.name,
            f.label,
            "BASE" if f.is_base else "PONTO",
            f.format_version or "não identificada",
            f.header_date or "não identificada",
            f.receiver_model or "não identificado",
            f.firmware_version or "não identificado",
            f.serial_id or "não identificado",
            f.zda_count,
            f.session_start_utc.isoformat() if f.session_start_utc else "n/d",
            f.session_end_utc.isoformat() if f.session_end_utc else "n/d",
        ])

    if result.single_zda_note:
        row = ws.max_row + 2
        ws.cell(
            row=row, column=1,
            value=(
                "Observação: apenas 1 arquivo apresentou sentenças ZDA. A duração de sessão só "
                "está disponível para esse arquivo (a base); isso é esperado, não um erro."
            ),
        )
        ws.cell(row=row, column=1).font = NOTE_FONT
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))

    _autosize(ws, len(headers))
    ws.freeze_panes = "A2"


def _sheet_auditoria(wb, result: ProcessingResult) -> None:
    ws = wb.create_sheet("Auditoria_Bytes")
    headers = [
        "Arquivo", "Ponto", "Offset início (hex)", "Bytes crus (hex)",
        "Lat_rad decodificado", "Lon_rad decodificado", "Alt_m decodificado",
    ]
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    for f in result.files:
        ws.append([
            f.path.name,
            f.label,
            "0x5A",
            f.audit_hex,
            f.lat_rad,
            f.lon_rad,
            f.alt_m,
        ])

    note_row = ws.max_row + 2
    ws.cell(
        row=note_row, column=1,
        value=(
            "Evidência bruta: bytes decodificados a partir do offset 0x5A (90 decimal) de cada "
            "arquivo — 8 bytes latitude (float64 LE, rad), 8 bytes longitude (float64 LE, rad), "
            "4 bytes altitude elipsoidal (float32 LE, m). Layout obtido por engenharia reversa; "
            "sem especificação oficial do fabricante."
        ),
    )
    ws.cell(row=note_row, column=1).font = NOTE_FONT
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=len(headers))
    ws.cell(row=note_row, column=1).alignment = Alignment(wrap_text=True)

    _autosize(ws, len(headers))
    ws.freeze_panes = "A2"


def _sheet_atualizacao_oficial(wb, result: ProcessingResult) -> None:
    ws = wb.create_sheet("Atualização_Coord_Oficial")

    ws.cell(row=1, column=1, value="Coordenada oficial da BASE (input do usuário)")
    ws.cell(row=1, column=1).font = Font(bold=True, size=12)

    ws.cell(row=3, column=1, value="Lat_oficial (graus)")
    ws.cell(row=3, column=2, value=result.base_geo.lat_deg)
    ws.cell(row=4, column=1, value="Lon_oficial (graus)")
    ws.cell(row=4, column=2, value=result.base_geo.lon_deg)
    ws.cell(row=5, column=1, value="Alt_oficial (m)")
    ws.cell(row=5, column=2, value=result.base_geo.alt_m)
    for r in (3, 4, 5):
        ws.cell(row=r, column=2).fill = PatternFill(
            start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"
        )

    ws.cell(row=7, column=1, value=(
        "Ao alterar B3/B4/B5 acima para a coordenada oficial (ex.: de processamento "
        "diferencial ou rede geodésica), a tabela abaixo recalcula automaticamente os "
        "pontos finais transladando os deltas calculados na aba Pontos_Resumo."
    ))
    ws.cell(row=7, column=1).font = NOTE_FONT
    ws.merge_cells(start_row=7, start_column=1, end_row=7, end_column=6)
    ws.cell(row=7, column=1).alignment = Alignment(wrap_text=True)

    headers_row = 9
    headers = ["Ponto", "Delta_E_m (ref. Pontos_Resumo)", "Delta_N_m (ref. Pontos_Resumo)",
               "Delta_h_m (ref. Pontos_Resumo)", "Lat_final_aprox (fórmula, graus)",
               "Lon_final_aprox (fórmula, graus)", "Alt_final (fórmula, m)"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=headers_row, column=c, value=h)
    _style_header(ws, headers_row, len(headers))

    # Fator aproximado de graus por metro (válido localmente, próximo à base).
    # 1 grau de latitude ~= 111320 m. 1 grau de longitude ~= 111320*cos(lat) m.
    ws.cell(row=headers_row + 1, column=1, value="_fator_lat_m_por_grau")
    ws.cell(row=headers_row + 1, column=2, value="=111320")
    ws.cell(row=headers_row + 2, column=1, value="_fator_lon_m_por_grau")
    ws.cell(row=headers_row + 2, column=2, value=f"=111320*COS(RADIANS($B$3))")

    start_data_row = headers_row + 4
    for i, p in enumerate(result.points_geo):
        row = start_data_row + i
        src_row = 3 + i  # linha correspondente em Pontos_Resumo (base=2, pontos a partir de 3)
        ws.cell(row=row, column=1, value=p.label)
        ws.cell(row=row, column=2, value=f"=Pontos_Resumo!M{src_row}")
        ws.cell(row=row, column=3, value=f"=Pontos_Resumo!N{src_row}")
        ws.cell(row=row, column=4, value=f"=Pontos_Resumo!O{src_row}")
        ws.cell(row=row, column=5, value=f"=$B$3+(C{row}/$B${headers_row+1})")
        ws.cell(row=row, column=6, value=f"=$B$4+(B{row}/$B${headers_row+2})")
        ws.cell(row=row, column=7, value=f"=$B$5+D{row}")

    note2 = start_data_row + len(result.points_geo) + 1
    ws.cell(
        row=note2, column=1,
        value=(
            "Nota: a translação lat/lon final usa um fator local metros->graus (aproximação "
            "plana válida para deltas pequenos, tipicamente até poucos km). Para redes extensas, "
            "recalcule com uma ferramenta geodésica completa."
        ),
    )
    ws.cell(row=note2, column=1).font = NOTE_FONT
    ws.merge_cells(start_row=note2, start_column=1, end_row=note2, end_column=6)
    ws.cell(row=note2, column=1).alignment = Alignment(wrap_text=True)

    _autosize(ws, 7)


def _sheet_metodologia(wb, result: ProcessingResult) -> None:
    ws = wb.create_sheet("Metodologia_Limitações")
    ws.column_dimensions["A"].width = 100

    linhas = [
        "METODOLOGIA",
        "",
        "1. Extração binária: cada arquivo .GNS é lido a partir do offset 0x5A (90 decimal), "
        "onde são extraídos 8 bytes de latitude (float64 little-endian, radianos), 8 bytes de "
        "longitude (float64 little-endian, radianos) e 4 bytes de altitude elipsoidal (float32 "
        "little-endian, metros).",
        "",
        "2. Detecção base/rover: o arquivo com maior número de sentenças ASCII 'ZDA' no corpo é "
        "tratado como BASE (estação contínua); os demais são tratados como pontos levantados "
        "(rovers), exceto quando o usuário indica manualmente qual arquivo é a base.",
        "",
        "3. Conversões: radianos -> graus decimais e DMS via fórmulas nativas (aba Pontos_Resumo) "
        "e via função Python equivalente (relatório docx/pdf). Projeção UTM via pyproj/PROJ "
        "(EPSG:4326 -> EPSG:326xx/327xx conforme zona/hemisfério detectados automaticamente).",
        "",
        "4. Deltas e distâncias: Delta E/N/h calculados em relação à BASE. Distância plana por "
        "Pitágoras sobre Delta E/N (projeção UTM). Distância geodésica e azimute calculados sobre "
        "o elipsoide WGS84 usando a solução de Karney (pyproj.Geod), não sobre a projeção plana.",
        "",
        "LIMITAÇÕES TÉCNICAS",
        "",
        "- O formato .GNS é proprietário e não documentado publicamente pelo fabricante. O layout "
        "de bytes usado nesta extração foi obtido por engenharia reversa e validação manual "
        "cruzada, não por especificação oficial.",
        "- Não há processamento diferencial (RINEX/GNSS baseline processing) neste pipeline: as "
        "coordenadas usadas são as gravadas no próprio arquivo .GNS pelo receptor, não uma solução "
        "pós-processada.",
        "- Não há informação de altura de antena, PDOP ou resíduos de processamento disponível "
        "nos arquivos de entrada; portanto essas métricas não constam neste pacote.",
        "- Os valores de UTM, distância geodésica e azimute na aba Pontos_Resumo são referência "
        "externa (calculados com pyproj/PROJ), não fórmulas nativas do Excel, pois reproduzir "
        "projeção cartográfica e geodésia elipsoidal em fórmulas de planilha não é praticável.",
    ]
    if result.single_zda_note:
        linhas.append(
            "- Apenas um arquivo do conjunto apresentou sentenças ZDA; a duração de sessão está "
            "disponível somente para esse arquivo (a base). Isso é esperado para levantamentos com "
            "um único ponto de referência contínuo, não um erro de processamento."
        )
    if result.cross_zone_points:
        pontos_str = ", ".join(p.label for p in result.cross_zone_points)
        linhas.append(
            f"- O(s) ponto(s) {pontos_str} está(ão) num fuso UTM diferente do fuso da BASE. Foi(foram) "
            f"reprojetado(s) automaticamente para o fuso da BASE antes do cálculo de Delta E/N e "
            f"distância plana (Pitágoras) — subtrair coordenadas de fusos UTM diferentes sem essa "
            f"reprojeção produziria um resultado incorreto. A distância geodésica e o azimute "
            f"(calculados diretamente sobre o elipsoide, via pyproj.Geod) não dependem de fuso UTM "
            f"e servem como referência cruzada independente."
        )

    for i, texto in enumerate(linhas, start=1):
        cell = ws.cell(row=i, column=1, value=texto)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if texto in ("METODOLOGIA", "LIMITAÇÕES TÉCNICAS"):
            cell.font = Font(bold=True, size=13)


def export_xlsx(result: ProcessingResult, output_path: Path) -> Path:
    wb = openpyxl.Workbook()
    _sheet_capa(wb, result)
    _sheet_pontos_resumo(wb, result)
    _sheet_metadados(wb, result)
    _sheet_auditoria(wb, result)
    _sheet_atualizacao_oficial(wb, result)
    _sheet_metodologia(wb, result)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    return output_path

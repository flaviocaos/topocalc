"""
docx_report.py
===============
Gera o relatório .docx. Todas as tabelas têm a largura de colunas calculada
programaticamente a partir da largura útil da página (page width - margens),
para garantir que nenhuma tabela ultrapasse a página.
"""

from __future__ import annotations

from pathlib import Path
from datetime import datetime

from docx import Document
from docx.shared import Cm, Pt, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from .pipeline import ProcessingResult
from .charts import plot_croqui, plot_distancias
from .assets import logo_png_path


def _usable_width_emu(document: Document) -> int:
    section = document.sections[0]
    return section.page_width - section.left_margin - section.right_margin


def _set_table_widths(table, usable_width_emu: int, weights: list[float]) -> None:
    """Define a largura de cada coluna proporcionalmente aos pesos
    informados, garantindo que a soma nunca ultrapasse a largura útil."""
    total_weight = sum(weights)
    col_widths = [int(usable_width_emu * (w / total_weight)) for w in weights]

    table.autofit = False
    tbl = table._tbl
    tblGrid = tbl.find(qn("w:tblGrid"))
    if tblGrid is None:
        tblGrid = OxmlElement("w:tblGrid")
        tbl.insert(0, tblGrid)
    for child in list(tblGrid):
        tblGrid.remove(child)
    for w in col_widths:
        gridCol = OxmlElement("w:gridCol")
        gridCol.set(qn("w:w"), str(Emu(w).twips))
        tblGrid.append(gridCol)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Emu(col_widths[idx])


def _shrink_font_for_wide_tables(table, ncols: int, base_pt: int = 10) -> int:
    """Reduz a fonte da tabela conforme o número de colunas, para caber
    dentro da largura calculada sem quebrar linhas de forma ilegível."""
    if ncols <= 5:
        return base_pt
    if ncols <= 8:
        return base_pt - 1
    return max(base_pt - 3, 6)


def _add_table(document, headers: list[str], rows: list[list[str]], weights: list[float] | None = None):
    ncols = len(headers)
    table = document.add_table(rows=1, cols=ncols)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    font_size = _shrink_font_for_wide_tables(table, ncols)

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(font_size)

    for row_data in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row_data):
            cells[i].text = str(val)
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(font_size)

    usable = _usable_width_emu(document)
    if weights is None:
        weights = [1.0] * ncols
    _set_table_widths(table, usable, weights)
    return table


def build_docx_report(result: ProcessingResult, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    document = Document()

    section = document.sections[0]
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    logo_path = logo_png_path()
    if logo_path is not None:
        usable_cm_header = (section.page_width - section.left_margin - section.right_margin) / 360000
        logo_p = document.add_paragraph()
        logo_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = logo_p.add_run()
        run.add_picture(str(logo_path), width=Cm(min(usable_cm_header * 0.55, 9)))

    document.add_heading("Relatório de Processamento GNSS — Pacote Geodésico", level=0)
    p = document.add_paragraph(f"Gerado em: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # 1. Identificação dos arquivos
    document.add_heading("1. Identificação dos Arquivos", level=1)
    headers = ["Arquivo", "Ponto", "Tipo", "Modelo", "Firmware", "Data", "ZDA"]
    rows = [
        [f.path.name, f.label, "BASE" if f.is_base else "PONTO",
         f.receiver_model or "n/d", f.firmware_version or "n/d",
         f.header_date or "n/d", str(f.zda_count)]
        for f in result.files
    ]
    _add_table(document, headers, rows, weights=[2.2, 1, 1, 1.2, 1, 1.2, 0.8])

    if result.single_zda_note:
        document.add_paragraph(
            "Observação: apenas um arquivo do conjunto apresentou sentenças ZDA. A duração de "
            "sessão está disponível somente para esse arquivo (a base). Isso é esperado, não "
            "indica erro no processamento."
        ).italic = True

    # 2. Metodologia
    document.add_heading("2. Metodologia", level=1)
    document.add_paragraph(
        "As coordenadas de cada ponto foram extraídas diretamente dos arquivos brutos .GNS, no "
        "offset 0x5A (90 decimal) do cabeçalho: 8 bytes de latitude (float64 little-endian, "
        "radianos), 8 bytes de longitude (float64 little-endian, radianos) e 4 bytes de altitude "
        "elipsoidal (float32 little-endian, metros)."
    )
    document.add_paragraph(
        "O arquivo tratado como BASE foi determinado pela maior contagem de sentenças ASCII "
        "'ZDA' no corpo do arquivo (estação contínua), salvo indicação manual do usuário."
    )
    document.add_paragraph(
        "As conversões de radianos para graus decimais e DMS, a projeção UTM (zona detectada "
        "automaticamente a partir da longitude) e os cálculos de distância geodésica/azimute "
        "elipsoidal (solução de Karney) foram realizados com as bibliotecas pyproj/PROJ, sem "
        "reimplementação manual de fórmulas de projeção."
    )
    if result.cross_zone_points:
        pontos_str = ", ".join(p.label for p in result.cross_zone_points)
        document.add_paragraph(
            f"Atenção — cruzamento de fuso UTM: o(s) ponto(s) {pontos_str} está(ão) naturalmente "
            f"em fuso UTM diferente do fuso {result.base_geo.native_utm_zone or result.base_geo.utm_zone}"
            f"{result.base_geo.utm_hemisphere} da BASE. Para permitir o cálculo direto de Delta "
            f"E/N e distância plana num único sistema de coordenadas, esse(s) ponto(s) foi(foram) "
            f"reprojetado(s) para o fuso da BASE antes do cálculo — prática padrão para redes "
            f"geodésicas locais. A distância geodésica e o azimute (calculados diretamente sobre "
            f"o elipsoide, independente de fuso UTM) não são afetados por esse ajuste e servem "
            f"como referência cruzada de conferência."
        ).italic = True

    # 3. Coordenadas consolidadas
    document.add_heading("3. Tabela Consolidada de Coordenadas", level=1)
    headers = ["Ponto", "Lat (graus)", "Lon (graus)", "Alt (m)", "Lat DMS", "Lon DMS", "UTM Zona", "Este (m)", "Norte (m)"]
    rows = []
    for g in result.all_geo():
        rows.append([
            g.label, f"{g.lat_deg:.8f}", f"{g.lon_deg:.8f}", f"{g.alt_m:.3f}",
            g.lat_dms, g.lon_dms, f"{g.utm_zone}{g.utm_hemisphere}",
            f"{g.utm_easting:.3f}", f"{g.utm_northing:.3f}",
        ])
    _add_table(document, headers, rows, weights=[0.8, 1.3, 1.3, 0.8, 1.4, 1.4, 0.8, 1, 1])

    # 4. Deltas / distâncias / azimutes
    document.add_heading("4. Deltas, Distâncias e Azimutes (em relação à BASE)", level=1)
    headers = ["Ponto", "Delta E (m)", "Delta N (m)", "Delta h (m)", "Dist. plana (m)", "Dist. geodésica (m)", "Azimute (°)"]
    rows = []
    for p in result.points_geo:
        rows.append([
            p.label, f"{p.delta_e:.3f}", f"{p.delta_n:.3f}", f"{p.delta_h:.3f}",
            f"{p.dist_plana_m:.3f}", f"{p.dist_geodesica_m:.3f}", f"{p.azimute_deg:.4f}",
        ])
    _add_table(document, headers, rows, weights=[0.8, 1, 1, 1, 1.2, 1.4, 1])

    # 5. Cálculo detalhado por ponto
    document.add_heading("5. Cálculo Detalhado por Ponto", level=1)
    base = result.base_geo
    for p in result.points_geo:
        document.add_heading(f"Ponto {p.label}", level=2)
        document.add_paragraph(
            f"Delta E = Este({p.label}) − Este(BASE) = {p.utm_easting:.3f} − {base.utm_easting:.3f} "
            f"= {p.delta_e:.3f} m"
        )
        document.add_paragraph(
            f"Delta N = Norte({p.label}) − Norte(BASE) = {p.utm_northing:.3f} − {base.utm_northing:.3f} "
            f"= {p.delta_n:.3f} m"
        )
        document.add_paragraph(
            f"Distância plana = √(ΔE² + ΔN²) = √({p.delta_e:.3f}² + {p.delta_n:.3f}²) = "
            f"{p.dist_plana_m:.3f} m"
        )
        document.add_paragraph(
            f"Distância geodésica (elipsoide WGS84, solução de Karney) = {p.dist_geodesica_m:.3f} m; "
            f"Azimute BASE→{p.label} = {p.azimute_deg:.4f}°"
        )

    # 6. Croqui e gráfico
    document.add_heading("6. Representação Gráfica", level=1)
    charts_dir = output_dir / "_charts"
    croqui_path = plot_croqui(result, charts_dir / "croqui.png")
    dist_path = plot_distancias(result, charts_dir / "distancias.png")

    usable_cm = (section.page_width - section.left_margin - section.right_margin) / 360000
    document.add_picture(str(croqui_path), width=Cm(min(usable_cm, 14)))
    document.add_paragraph("Figura 1 — Croqui de distribuição espacial dos pontos.").italic = True
    document.add_picture(str(dist_path), width=Cm(min(usable_cm, 16)))
    document.add_paragraph("Figura 2 — Distância de cada ponto em relação à BASE.").italic = True

    # 7. Limitações técnicas
    document.add_heading("7. Limitações Técnicas", level=1)
    limitacoes = [
        "Formato .GNS proprietário e não documentado oficialmente pelo fabricante; o layout de "
        "bytes usado nesta extração foi obtido por engenharia reversa e validação manual, não por "
        "especificação oficial.",
        "Não há processamento diferencial (RINEX / GNSS baseline processing): as coordenadas são "
        "as gravadas pelo receptor no próprio arquivo .GNS, não uma solução pós-processada.",
        "Não há informação de altura de antena, PDOP ou resíduos de processamento nos arquivos de "
        "entrada; portanto essas métricas não constam neste relatório.",
    ]
    if result.single_zda_note:
        limitacoes.append(
            "Apenas um arquivo apresentou sentenças ZDA; a duração de sessão está disponível "
            "somente para esse arquivo (a base) — comportamento esperado, não um erro."
        )
    if result.cross_zone_points:
        pontos_str = ", ".join(p.label for p in result.cross_zone_points)
        limitacoes.append(
            f"O(s) ponto(s) {pontos_str} está(ão) num fuso UTM diferente do fuso da BASE e foi(foram) "
            f"reprojetado(s) para o fuso da BASE para permitir o cálculo de deltas/distância plana "
            f"(ver seção 2, Metodologia). Ao se afastar do meridiano central do fuso da BASE, a "
            f"distorção de escala da projeção UTM aumenta — para redes muito extensas ou próximas "
            f"de um limite de fuso, prefira a distância geodésica (coluna 'Dist. geodésica') como "
            f"referência principal, não a distância plana."
        )
    for item in limitacoes:
        document.add_paragraph(item, style="List Bullet")

    # 8. Auditoria de bytes
    document.add_heading("8. Auditoria — Bytes Brutos Decodificados", level=1)
    document.add_paragraph(
        "Evidência bruta dos bytes lidos a partir do offset 0x5A de cada arquivo, para permitir "
        "conferência independente da extração."
    )
    headers = ["Arquivo", "Offset", "Bytes (hex)"]
    rows = [[f.path.name, "0x5A", f.audit_hex] for f in result.files]
    _add_table(document, headers, rows, weights=[1.4, 0.7, 4])

    out_path = output_dir / "relatorio_geodesico.docx"
    document.save(str(out_path))
    return out_path

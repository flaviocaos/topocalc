"""Planilha .xlsx com todos os resultados do processamento."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .common import SOLUTION_LABEL, baseline_rows, dms, fmt_day, fmt_dt

TEAL = "1E8A93"
LIGHT = "E3F1F2"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def _sheet(wb, title, headers, rows, widths=None, numfmt=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=TEAL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for i, c in enumerate(row):
            c.border = BORDER
            if numfmt and i < len(numfmt) and numfmt[i] and isinstance(c.value, (int, float)):
                c.number_format = numfmt[i]
    for i, h in enumerate(headers, 1):
        w = (widths[i - 1] if widths else None) or max(10, min(40, len(str(h)) + 2))
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"
    return ws


def export_xlsx(p: dict, path: Path) -> Path:
    res = p["results"]
    wb = Workbook()
    wb.remove(wb.active)
    pts = res["points"]
    z = f"{pts[0]['zone']}{pts[0]['hemisphere']}" if pts else ""

    _sheet(wb, "Coordenadas",
           ["Ponto", f"E (m) UTM {z}", f"N (m) UTM {z}", "h elipsoidal (m)", "N geoidal (m)",
            "H ortométrica (m)", "σE (m)", "σN (m)", "σh (m)", "Tipo", "Vetores fixos", "Vetores flutuantes"],
           [[pt["name"], pt["E"], pt["N"], pt["h"], pt["N_geoid"], pt["H"], pt["sE"], pt["sN"], pt["sU"],
             "Controle (fixo)" if pt["fixed"] else ("Controle" if pt["type"] == "controle" else "Ajustado"),
             pt["vectors"]["fixa"], pt["vectors"]["flutuante"]] for pt in pts],
           [12, 15, 15, 15, 13, 15, 9, 9, 9, 16, 12, 14],
           [None, "0.000", "0.000", "0.000", "0.000", "0.000", "0.000", "0.000", "0.000"])

    extra = res.get("extra_crs") or {}
    if extra:
        hdr, keys = ["Ponto"], []
        for k, label in sorted(extra.items()):
            if label.startswith("indisponível"):
                continue
            keys.append(k)
            hdr += [f"{label} — E (m)", "N (m)"] + (["H (m)"] if k == "ltp" else [])
        rows = []
        for pt in pts:
            row = [pt["name"]]
            for k in keys:
                v = (pt.get("extra") or {}).get(k)
                row += [v[0] if v else None, v[1] if v else None] + ([pt["H"]] if k == "ltp" else [])
            rows.append(row)
        _sheet(wb, "Outros sistemas", hdr, rows, [12] + [24] * (len(hdr) - 1),
               [None] + ["0.000"] * (len(hdr) - 1))

    _sheet(wb, "Geodésicas",
           ["Ponto", "Latitude (°)", "Longitude (°)", "Latitude (GMS)", "Longitude (GMS)", "h (m)",
            "Convergência meridiana (°)", "Fator de escala UTM"],
           [[pt["name"], pt["lat"], pt["lon"], dms(pt["lat"], "lat"), dms(pt["lon"], "lon"), pt["h"],
             pt["convergence"], pt["scale"]] for pt in pts],
           [12, 15, 15, 20, 20, 12, 16, 14],
           [None, "0.000000000", "0.000000000", None, None, "0.000", "0.000000", "0.00000000"])

    rows = []
    for r in baseline_rows(p):
        a = r.get("adj") or {}
        rows.append([r["day"], r["base"], r["rover"], r["length_m"], r["duration_min"],
                     SOLUTION_LABEL.get(r["solution"], r["solution"]), r["ratio"] or None,
                     *(r["denu"] or [None] * 3), r["sigma_h_m"], r["sigma_v_m"],
                     *((a.get("v_enu_mm") or [None] * 3)), a.get("w_max"),
                     "EXCLUÍDA" if r["excluded"] else (a.get("flag") or ""),
                     "; ".join(r["notes"]), r["method"]])
    _sheet(wb, "Linhas de base",
           ["Dia", "Base", "Rover", "Comprimento (m)", "Sessão comum (min)", "Solução", "Ratio",
            "ΔE (m)", "ΔN (m)", "ΔU (m)", "σ horiz. a priori (m)", "σ vert. a priori (m)",
            "Resíduo E (mm)", "Resíduo N (mm)", "Resíduo U (mm)", "|w| máx", "Situação", "Observações",
            "Método"], rows,
           [11, 9, 9, 13, 11, 11, 8, 13, 13, 11, 11, 11, 10, 10, 10, 8, 11, 45, 50],
           [None, None, None, "0.000", "0", None, "0.0", "0.0000", "0.0000", "0.0000", "0.000", "0.000",
            "0.0", "0.0", "0.0", "0.00"])

    a = res["adjustment"]
    _sheet(wb, "Ajustamento", ["Item", "Valor"], [
        ["Vetores no ajustamento", a["n_vectors"]],
        ["Vetores excluídos", a["n_excluded"]],
        ["Graus de liberdade", a["dof"]],
        ["Variância a posteriori (σ0²)", a["sigma0_sq"]],
        ["Valor crítico (F 95 %)", a["f_crit"]],
        ["Teste global", "Aprovado" if a["global_test_passed"] else "Reprovado"],
        ["Pontos sem ligação com o controle", ", ".join(a["unconnected"]) or "nenhum"],
        ["Referencial", res.get("frame") or "WGS 84"],
        ["Órbitas", "; ".join(f"{d}: {o}" for d, o in sorted((res.get("orbits") or {}).items()))],
        ["Calibração de antena", ", ".join(res.get("antennas") or []) or "não aplicada"],
        ["Geoide", res["geoid"]["model"]],
        ["Ajuste do geoide ao controle (m)", res["geoid"]["offset"]],
    ], [38, 30])

    _sheet(wb, "Fechamentos",
           ["Triângulo", "Dias", "Perímetro (m)", "Fechamento (mm)", "E (mm)", "N (mm)", "U (mm)", "ppm"],
           [[" – ".join(l["points"]), ", ".join(fmt_day(d) for d in l["days"]), l["perimeter_m"],
             l["misclosure_mm"], *(l.get("mis_enu_mm") or [None] * 3), l["ppm"]] for l in res["loops"]],
           [28, 24, 13, 13, 9, 9, 9, 8],
           [None, None, "0.0", "0.0", "0.0", "0.0", "0.0", "0.00"])

    _sheet(wb, "Arquivos",
           ["Arquivo", "Ponto", "Dia", "Início (GPS)", "Fim (GPS)", "Épocas", "Intervalo (s)", "Receptor",
            "Série", "Formato interno", "Altura da antena (m)", "Incluído"],
           [[o["file"], o["name"], fmt_day(o["day"]), fmt_dt(o["first"]), fmt_dt(o["last"]), o["n_epochs"],
             o["interval_s"], o["receiver"], o["serial"], o["format"], o["antenna_height"],
             "sim" if o.get("include", True) else "não"] for o in p["occupations"]],
           [16, 10, 11, 16, 16, 8, 10, 12, 13, 14, 12, 9])

    _sheet(wb, "Controle",
           ["Marco", "Latitude", "Longitude", "h (m)", "H (m)", "σ horiz. (m)", "σ vert. (m)", "Observações"],
           [[c["name"], dms(c["lat"], "lat"), dms(c["lon"], "lon"), c["h"], c.get("H"),
             c.get("sigma_h") or "fixo", c.get("sigma_v") or "fixo", "; ".join(c.get("notes", []))]
            for c in p["controls"]],
           [10, 20, 20, 12, 12, 11, 11, 40])

    wb.save(path)
    return path

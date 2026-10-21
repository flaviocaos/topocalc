"""
docx_report.py
==============
Relatório técnico de pós-processamento GNSS (.docx). O PDF é gerado a partir
deste mesmo arquivo, para que os dois nunca divirjam.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Emu, Pt, RGBColor  # noqa: E402

from .. import __version__  # noqa: E402
from ..assets import logo_png_path  # noqa: E402
from ..gnss.baseline import APRIORI, SHORT_KM, TOL_FIX_FLOAT, TOL_FWD_BWD  # noqa: E402
from .common import SOLUTION_LABEL, baseline_rows, dms, fmt_day, fmt_hm  # noqa: E402

TEAL = RGBColor(0x1E, 0x8A, 0x93)
INK = RGBColor(0x3A, 0x3A, 0x39)


# ------------------------------------------------------------ utilidades

def _shade(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _usable(doc) -> int:
    s = doc.sections[-1]
    return s.page_width - s.left_margin - s.right_margin


def _table(doc, headers, rows, weights=None, font=8.5, highlight=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(font)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _shade(c, "1E8A93")
    for k, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run("" if v is None else str(v))
            r.font.size = Pt(font)
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if highlight and highlight(k):
                _shade(cells[i], highlight(k))
    weights = weights or [1] * len(headers)
    tot = sum(weights)
    widths = [int(_usable(doc) * w / tot) for w in weights]
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(Emu(w).twips))
    for row in t.rows:
        for i, c in enumerate(row.cells):
            c.width = Emu(widths[i])
    doc.add_paragraph()
    return t


def _para(doc, text, italic=False, size=10, bold=False, color=None, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    r.bold = bold
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    if align:
        p.alignment = align
    return p


def _bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text).font.size = Pt(10)
    return p


def _heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = TEAL if level == 1 else INK
    return h


def _croqui(p: dict, path: Path) -> Path:
    pts = {pt["name"]: pt for pt in p["results"]["points"]}
    fig, ax = plt.subplots(figsize=(7.0, 6.2), dpi=160)
    colors = {"fixa": "#2E9E44", "flutuante": "#F0A020"}
    seen = set()
    for r in baseline_rows(p):
        a, b = pts.get(r["base"]), pts.get(r["rover"])
        if not a or not b:
            continue
        if r["excluded"]:
            col, ls, lab = "#C0392B", ":", "excluída"
        else:
            col, ls, lab = colors.get(r["solution"], "gray"), "-", r["solution"]
        ax.plot([a["E"], b["E"]], [a["N"], b["N"]], ls, color=col, lw=1.6, zorder=1,
                label=None if lab in seen else f"vetor {lab}")
        seen.add(lab)
    for pt in pts.values():
        ctrl = pt["type"] == "controle"
        ax.scatter([pt["E"]], [pt["N"]], marker="^" if ctrl else "o", s=110 if ctrl else 55,
                   c="#C0392B" if ctrl else "#1E8A93", edgecolors="white", zorder=3)
        ax.annotate(pt["name"], (pt["E"], pt["N"]), textcoords="offset points", xytext=(6, 6),
                    fontsize=8.5, fontweight="bold" if ctrl else None, color="#3A3A39")
    z = next(iter(pts.values()))
    ax.set_xlabel(f"E (m) — UTM {z['zone']}{z['hemisphere']}")
    ax.set_ylabel("N (m)")
    ax.ticklabel_format(useOffset=False, style="plain")
    ax.tick_params(labelsize=7)
    ax.grid(True, ls=":", alpha=0.5)
    ax.set_aspect("equal", adjustable="datalim")
    ax.legend(fontsize=8, loc="best")
    ax.set_title("Croqui da rede GNSS processada", fontsize=10, color="#3A3A39")
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)
    return path


# ------------------------------------------------------------ relatório

def build_docx(p: dict, path: Path, workdir: Path) -> Path:
    res = p["results"]
    adj = res["adjustment"]
    pts = res["points"]
    bls = baseline_rows(p)
    doc = Document()
    sec = doc.sections[0]
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Cm(1.8))
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.6)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)

    logo = logo_png_path()
    if logo:
        doc.add_paragraph().add_run().add_picture(str(logo), width=Cm(6))
    t = doc.add_heading("Relatório de Pós-Processamento GNSS", level=0)
    for r in t.runs:
        r.font.color.rgb = INK
    _para(doc, p.get("title") or "", size=13, bold=True, color=TEAL)
    days = sorted({o["day"] for o in p["occupations"] if o.get("include", True)})
    _para(doc, "Rastreio: " + ", ".join(fmt_day(d) for d in days)
          + f"   •   Processado em {datetime.now():%d/%m/%Y %H:%M}   •   TopoCalc {__version__}",
          size=9, color=INK)

    # 1. Resumo
    _heading(doc, "1. Resumo")
    nfix = sum(1 for b in bls if b["solution"] == "fixa" and not b["excluded"])
    nflo = sum(1 for b in bls if b["solution"] == "flutuante" and not b["excluded"])
    nexc = sum(1 for b in bls if b["excluded"])
    nfail = sum(1 for b in bls if b["solution"] == "falhou")
    ctrl_names = ", ".join(c["name"] for c in p["controls"])
    _para(doc, (
        f"Foram processados {len(p['occupations'])} arquivos brutos (.GNS) de {len(days)} dia(s) de rastreio, "
        f"formando {len(bls)} linhas de base: {nfix} com ambiguidades fixas, {nflo} flutuantes"
        + (f", {nfail} sem solução" if nfail else "") + (f" e {nexc} excluída(s) do ajustamento" if nexc else "")
        + f". A rede foi ajustada por mínimos quadrados a partir do controle {ctrl_names}. "
        f"O teste global do ajustamento foi {'APROVADO' if adj['global_test_passed'] else 'REPROVADO'} "
        f"(σ0² = {adj['sigma0_sq']:.2f}; valor crítico {adj['f_crit']:.2f}; {adj['dof']} graus de liberdade)."
        if adj["dof"] > 0 else
        f"formando {len(bls)} linhas de base. A rede não tem redundância: cada ponto foi determinado por um "
        "único vetor, sem verificação independente."))
    if adj["unconnected"]:
        _para(doc, "ATENÇÃO: os pontos " + ", ".join(adj["unconnected"])
              + " não têm ligação com o controle e não receberam coordenadas.", bold=True,
              color=RGBColor(0xC0, 0x39, 0x2B))

    # 2. Coordenadas finais
    _heading(doc, "2. Coordenadas finais")
    z = f"{pts[0]['zone']}{pts[0]['hemisphere']}"
    _para(doc, f"Referencial: {res.get('frame') or 'WGS 84'} (o mesmo das coordenadas do controle), projeção "
               f"UTM fuso {z}. Altitude ortométrica H = h − N, com N do modelo {res['geoid']['model']}.",
          size=9, italic=True)
    _table(doc, ["Ponto", "E (m)", "N (m)", "h (m)", "N geoidal (m)", "H (m)", "σE (m)", "σN (m)",
                 "σh (m)", "Tipo"],
           [[pt["name"], f"{pt['E']:.3f}", f"{pt['N']:.3f}", f"{pt['h']:.3f}", f"{pt['N_geoid']:.3f}",
             f"{pt['H']:.3f}", f"{pt['sE']:.3f}", f"{pt['sN']:.3f}", f"{pt['sU']:.3f}",
             "Controle" if pt["type"] == "controle" else "Ajustado"] for pt in pts],
           [1, 1.5, 1.6, 1.2, 1.1, 1.2, 0.85, 0.85, 0.85, 1.1])
    _heading(doc, "2.1 Coordenadas geodésicas", 2)
    _table(doc, ["Ponto", "Latitude", "Longitude", "h (m)", "Convergência (°)", "Fator de escala"],
           [[pt["name"], dms(pt["lat"], "lat"), dms(pt["lon"], "lon"), f"{pt['h']:.3f}",
             f"{pt['convergence']:.6f}", f"{pt['scale']:.8f}"] for pt in pts],
           [1, 2, 2, 1.1, 1.3, 1.3])

    extra = res.get("extra_crs") or {}
    if extra:
        _heading(doc, "2.2 Outros sistemas de coordenadas", 2)
        for key, label in sorted(extra.items()):
            if label.startswith("indisponível"):
                _para(doc, f"{key}: {label}", size=9, italic=True)
                continue
            _para(doc, label, size=9, bold=True)
            if key == "ltp":
                _table(doc, ["Ponto", "E local (m)", "N local (m)", "H (m)"],
                       [[pt["name"], f"{pt['extra']['ltp'][0]:.3f}", f"{pt['extra']['ltp'][1]:.3f}",
                         f"{pt['H']:.3f}"] for pt in pts if "ltp" in pt.get("extra", {})], [1, 1.5, 1.5, 1.2])
                _para(doc, "Plano tangente ao elipsoide na origem: distâncias e áreas no plano correspondem às do "
                           "terreno perto da altitude da origem. Não é a formulação completa da NBR 14166.",
                      size=8.5, italic=True)
            else:
                _table(doc, ["Ponto", "E (m)", "N (m)"],
                       [[pt["name"], f"{pt['extra'][key][0]:.3f}", f"{pt['extra'][key][1]:.3f}"]
                        for pt in pts if key in pt.get("extra", {})], [1, 1.5, 1.5])
                _para(doc, "Transformação oficial EPSG \"Camacupa 2015 to WGS 84 (11)\", precisão declarada de "
                           "±3 m. Serve para compatibilidade com a cartografia angolana; não substitui "
                           "parâmetros locais determinados em marcos da rede oficial.", size=8.5, italic=True)

    croqui = _croqui(p, workdir / "croqui.png")
    doc.add_picture(str(croqui), width=Cm(15.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 3. Controle
    _heading(doc, "3. Controle e altimetria")
    _table(doc, ["Marco", "Latitude", "Longitude", "h (m)", "H (m)", "Injunção"],
           [[c["name"], dms(c["lat"], "lat"), dms(c["lon"], "lon"), f"{c['h']:.4f}",
             "" if c.get("H") is None else f"{c['H']:.4f}",
             "fixo" if not (c.get("sigma_h") or c.get("sigma_v"))
             else f"σH {c['sigma_h']:.3f} m / σV {c['sigma_v']:.3f} m"] for c in p["controls"]],
           [1, 2, 2, 1.2, 1.2, 1.6])
    for c in p["controls"]:
        if c.get("ppp"):
            _para(doc, f"{c['name']}: " + " ".join(c.get("notes", [])), size=9)
    g = res["geoid"]
    if g["controls"]:
        txt = "; ".join(f"{c['name']}: N observado {c['N_obs']:.3f} m, EGM96 {c['N_egm96']:.3f} m"
                        for c in g["controls"])
        _para(doc, f"Ondulação geoidal: EGM96 (grade 15') deslocado de {g['offset']:+.3f} m para coincidir com "
                   f"a ondulação observada no controle ({txt}). A diferença entre os pontos segue a forma do "
                   "EGM96; em distâncias de dezenas de quilômetros a incerteza relativa do modelo é de alguns "
                   "centímetros a decímetros.", size=9)
    else:
        _para(doc, "Ondulação geoidal: EGM96 (grade 15'), sem ajuste local — nenhum controle com altitude "
                   "ortométrica informada. As altitudes H têm a incerteza absoluta do EGM96 (decímetros a metro).",
              size=9)

    # 4. Dados
    _heading(doc, "4. Dados de campo")
    _table(doc, ["Arquivo", "Ponto", "Dia", "Início", "Fim", "Receptor", "Formato", "Alt. antena (m)"],
           [[o["file"], o["name"], fmt_day(o["day"]), fmt_hm(o["first"]), fmt_hm(o["last"]),
             o["receiver"] or "", "NovAtel" if o["format"] == "novatel" else "Trimble RT27",
             f"{o['antenna_height']:.3f}"] for o in p["occupations"] if o.get("include", True)],
           [1.6, 0.9, 1.0, 0.7, 0.7, 0.9, 1.1, 1.0])
    _para(doc, "Horários em tempo GPS. Altura da antena: vertical do marco ao ponto de referência da antena "
               "(ARP), conforme registrada no receptor ou corrigida na revisão.", size=8.5, italic=True)

    # 5. Linhas de base
    _heading(doc, "5. Linhas de base")
    _table(doc, ["Dia", "Base → Rover", "Compr. (km)", "Sessão (min)", "Solução", "Ratio", "Ida−volta (mm)",
                 "σH/σV a priori (mm)", "Situação"],
           [[fmt_day(r["day"]), f"{r['base']} → {r['rover']}", f"{r['length_m'] / 1000:.3f}",
             f"{r['duration_min']:.0f}", SOLUTION_LABEL.get(r["solution"], r["solution"]),
             f"{r['ratio']:.1f}" if r["ratio"] else "",
             "/".join(f"{v:.0f}" for v in r["fwd_bwd_mm"]) if r["fwd_bwd_mm"] else "",
             f"{r['sigma_h_m'] * 1000:.0f}/{r['sigma_v_m'] * 1000:.0f}",
             "EXCLUÍDA" if r["excluded"] else ((r.get("adj") or {}).get("flag") or "")] for r in bls],
           [1, 1.8, 0.9, 0.8, 0.9, 0.6, 1.2, 1.1, 0.9],
           highlight=lambda k: "FDECEA" if bls[k]["excluded"] or bls[k]["solution"] == "falhou"
           else ("FFF4E0" if bls[k]["solution"] == "flutuante" else None))
    notes = [(f"{r['base']} → {r['rover']}", n) for r in bls for n in r["notes"]]
    if notes:
        _heading(doc, "5.1 Observações do processamento", 2)
        for who, n in notes:
            _bullet(doc, f"{who}: {n}")

    # 6. Ajustamento
    _heading(doc, "6. Ajustamento em rede")
    _table(doc, ["Item", "Valor"], [
        ["Vetores ajustados / excluídos", f"{adj['n_vectors']} / {adj['n_excluded']}"],
        ["Graus de liberdade", str(adj["dof"])],
        ["Variância a posteriori σ0²", f"{adj['sigma0_sq']:.3f}" if adj["dof"] else "—"],
        ["Valor crítico (F, 95 %)", f"{adj['f_crit']:.3f}" if adj["dof"] else "—"],
        ["Teste global", ("Aprovado" if adj["global_test_passed"] else "Reprovado") if adj["dof"] else "—"],
    ], [2, 1.5])
    if adj["dof"]:
        _para(doc, "Resíduos dos vetores (componentes locais E, N, U) e estatística w de Baarda. "
                   "|w| > 1,96 merece atenção; |w| > 3,29 indica erro grosseiro provável — o vetor deve "
                   "ser excluído e a rede reajustada.", size=9)
        obs = res["observations"]
        _table(doc, ["Vetor", "vE (mm)", "vN (mm)", "vU (mm)", "wE", "wN", "wU", "Redund.", "Situação"],
               [[f"{o['base']} → {o['rover']}", *(f"{v:.1f}" for v in o["v_enu_mm"]),
                 *(f"{w:.2f}" for w in o["w"]), f"{o['redundancy']:.2f}", o["flag"]] for o in obs],
               [1.8, 0.8, 0.8, 0.8, 0.7, 0.7, 0.7, 0.8, 1],
               highlight=lambda k: "FDECEA" if obs[k]["flag"] == "rejeitar"
               else ("FFF4E0" if obs[k]["flag"] == "atenção" else None))
    if res["loops"]:
        _heading(doc, "6.1 Fechamento de triângulos", 2)
        _para(doc, "Soma dos três vetores de cada triângulo (deveria ser zero). Independe do controle e do "
                   "ajustamento; triângulos com vetores de dias diferentes testam também a reocupação.", size=9)
        loops = sorted(res["loops"], key=lambda l: -l["ppm"])[:30]
        _table(doc, ["Triângulo", "Perímetro (km)", "E (mm)", "N (mm)", "U (mm)", "Total (mm)", "ppm"],
               [[" – ".join(l["points"]), f"{l['perimeter_m'] / 1000:.2f}",
                 *(f"{v:.1f}" for v in (l.get("mis_enu_mm") or [0, 0, 0])), f"{l['misclosure_mm']:.1f}",
                 f"{l['ppm']:.2f}"] for l in loops],
               [2.4, 1, 0.8, 0.8, 0.8, 0.9, 0.7])

    # 7. Metodologia
    _heading(doc, "7. Metodologia")
    (ah, bh), (av, bv) = APRIORI["fixa"]
    (fh, gh), (fv, gv) = APRIORI["flutuante"]
    for t in [
        "Conversão: os arquivos .GNS (Hi-Target) foram convertidos para RINEX 3.04. Receptores com placa "
        "NovAtel: observações e efemérides GLONASS/Galileo pelo RTKLIB (convbin); efemérides GPS "
        "(GPSEPHEM) por decodificador próprio. Receptores com placa Trimble: observações do registro RT27 "
        "por decodificador próprio (validado contra o Leica Infinity) e efemérides GPS pelo RTKLIB (RT17).",
        "Processamento: RTKLIB demo5 (rnx2rtkp), modo estático, GPS + Galileo, máscara de elevação 15°, "
        "órbitas: " + ("; ".join(f"{fmt_day(d)} {o}" for d, o in sorted((res.get("orbits") or {}).items()))
                       or "transmitidas") + ". "
        f"Linhas com menos de {SHORT_KM:.0f} km: L1+L2, ionosfera do modelo transmitido, ambiguidades "
        "resolvidas para frente e para trás. Linhas longas: ionosfera e troposfera estimadas; solução "
        "flutuante e tentativa de fixação nos dois sentidos.",
        f"Validação da fixação (linhas longas): a solução fixa só é aceita se ida e volta estiverem fixas "
        f"(ratio ≥ 3) e concordarem em até {TOL_FWD_BWD[0] * 1000:.0f} mm (horizontal) / "
        f"{TOL_FWD_BWD[1] * 1000:.0f} mm (vertical), e se concordar com a solução flutuante em até "
        f"{TOL_FIX_FLOAT[0] * 1000:.0f} / {TOL_FIX_FLOAT[1] * 1000:.0f} mm. Caso contrário usa-se a "
        "flutuante. Isso evita aceitar ambiguidades fixadas erradas, comuns em linhas longas.",
        f"Pesos: precisão a priori de {ah * 1000:.0f} mm + {bh} ppm (horizontal) e {av * 1000:.0f} mm + {bv} ppm "
        f"(vertical) para vetores fixos; {fh * 1000:.0f} mm + {gh} ppm e {fv * 1000:.0f} mm + {gv} ppm para "
        "flutuantes. As precisões finais são as do ajustamento, ampliadas por σ0² quando maior que 1.",
        "Ajustamento: mínimos quadrados 3D em coordenadas cartesianas, com o(s) controle(s) fixo(s), teste "
        "global da variância a posteriori e teste w (data snooping) por componente.",
    ]:
        _bullet(doc, t)

    _heading(doc, "8. Limitações")
    lim = []
    if res.get("antennas"):
        _bullet(doc, "Calibração de antena (centro de fase) aplicada: " + ", ".join(res["antennas"]) + ".")
        if any(not o.get("antenna_type") for o in p["occupations"] if o.get("include", True)):
            lim.append("Alguns arquivos ficaram sem modelo de antena: neles a calibração não foi aplicada.")
    else:
        lim.append("Calibração de antena (centro de fase) não aplicada: nenhum modelo de antena foi escolhido. "
                   "Com antenas de modelos diferentes nas duas pontas, isso pode causar diferenças de alguns "
                   "centímetros na altura.")
    if any("transmitidas" in o for o in (res.get("orbits") or {}).values()):
        lim.append("Nos dias processados com efemérides transmitidas, o erro de órbita contribui com cerca de "
                   "0,1–0,5 ppm do comprimento da linha. Para linhas acima de ~20 km, reprocesse quando as "
                   "órbitas precisas estiverem disponíveis.")
    if nflo:
        lim.append(f"{nflo} vetor(es) com solução flutuante: precisão de centímetros, não de milímetros. "
                   "Pontos determinados só por vetores flutuantes devem ser tratados com essa precisão.")
    if not adj["dof"]:
        lim.append("A rede não tem redundância; um erro em qualquer vetor passa despercebido.")
    for t in lim:
        _bullet(doc, t)

    doc.save(path)
    return path
